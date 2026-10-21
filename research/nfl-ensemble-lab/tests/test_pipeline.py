"""End-to-end smoke test on the synthetic league: training, walk-forward
separation, stacking, selection-before-test, weekly predictions, immutable
snapshots, Excel exports and Streamlit startup."""
import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from lab import ensemble as ens
from lab import training, weekly
from lab.config import VALIDATION_SEASONS, WALK_FORWARD_START


def test_walk_forward_is_out_of_sample(trained):
    oof = trained["leaderboard"], training.load_run()["oof"]
    oof = oof[1]
    assert oof["season"].min() == WALK_FORWARD_START
    for s, part in oof.groupby("season"):
        lo, hi = map(int, part["train_seasons"].iloc[0].split("-"))
        assert hi == s - 1, "season s must be predicted by models trained only on seasons < s"
    for k in ens.MODEL_KEYS:
        assert oof[f"p_{k}"].between(0, 1).all()


def test_stacking_uses_only_earlier_seasons(trained):
    oof = training.load_run()["oof"]
    assert oof.loc[oof["season"] == WALK_FORWARD_START, "p_stacking"].isna().all()
    s = 2022
    pcols = [f"p_{k}" for k in ens.MODEL_KEYS]
    hist = oof[(oof["season"] < s) & oof["home_win"].notna()]
    expect = ens.stack_predict(ens.fit_stacker(hist[pcols].to_numpy(), hist["home_win"].to_numpy()),
                               oof.loc[oof["season"] == s, pcols].to_numpy())
    np.testing.assert_allclose(oof.loc[oof["season"] == s, "p_stacking"], expect)


def test_champion_selected_on_validation_only(trained):
    run = training.load_run()
    sel = run["selection"]
    assert sel["validation_seasons"] == list(VALIDATION_SEASONS)
    tie = sel["tie_analysis"]
    assert sel["champion"] in tie["tied_within_1se"] and tie["best_on_validation"] in tie["tied_within_1se"]
    assert max(tie["history_seasons"]) <= max(VALIDATION_SEASONS)
    assert "test" in set(run["leaderboard"]["split"])
    # all six winner models and margin models produced out-of-sample results
    lb = run["leaderboard"]
    assert set(ens.MODEL_KEYS) <= set(lb[lb["split"] == "test"]["method"])
    assert {"majority_vote", "equal_average", "weighted_average", "stacking", "market"} <= set(lb["method"])
    mg = run["margins"]
    assert {"ensemble", "spread"} <= set(mg["method"]) and mg["mae"].notna().all()


@pytest.fixture(scope="module")
def disc_run(trained):
    from lab import discovery
    cfg = {"rule_trees": 5, "pair_top_features": 6, "symbolic_top_features": 5, "gbm_top_features": 15, "n_clusters": 3,
           "min_n_discovery": 15, "min_n_validation": 10}
    return discovery.run_experiment(cfg, progress=lambda m: None)


@pytest.fixture(scope="module")
def week_run(trained, disc_run):
    return weekly.run_weekly(season=2026, week=6, refresh=False, progress=lambda m: None)


def test_discovery_on_pipeline_outputs(disc_run, week_run):
    s = disc_run["summary"]
    assert set(s["baselines"]) == {"ensemble", "market"}
    for info in s["baselines"].values():
        assert info["origins"]["primary"]["n_tested"] > 100
    d = week_run["discovery"]
    assert d["experiment"] == s["experiment"] and len(d["rows"]) == 4
    t = week_run["table"]
    assert {"p_home_adj_patterns", "p_home_adj_gbm", "adj_patterns"} <= set(t.columns)
    assert t["p_home_adj_gbm"].between(0, 1).all()
    from lab.reports import discovery_workbook
    wb = load_workbook(discovery_workbook(weekly_disc=d))
    assert wb.sheetnames == ["Plain-English Summary", "Discovery Leaderboard", "Discovered Rules", "Feature Interactions",
                             "Historical Matches", "Expected vs Observed", "Validation Results", "Rejected Hypotheses",
                             "Upcoming Upset Candidates", "Methodology"]


def test_weekly_predictions_and_debate(week_run):
    t, debates = week_run["table"], week_run["debates"]
    assert len(t) == 4 and len(debates) == 4
    assert ((t["home_votes"] + t["away_votes"]) == 6).all()
    for d in debates:
        assert len(d["round1"]) == 6 and set(d["round2"]) == set(ens.MODEL_KEYS)
        assert d["round2"]["decision_tree"]["rules"]
        for e in d["round2"].values():
            assert len(e["top"]) > 0 and e["scope"].startswith("local")
        assert "vote_vs_ensemble_disagree" in d["round5"]


def test_explanations_reconstruct_linear_models(week_run, trained):
    """Coefficient attributions must sum to the model's logit."""
    bundle = training.load_bundle()
    fs = pd.read_parquet(training.load_feature_store.__globals__["FEATURE_STORE"])
    g = fs[(fs["season"] == 2026) & (fs["week"] == 6)].iloc[[0]]
    from lab import explain
    X = g[bundle["features"]].astype(float).to_numpy()
    e = explain.explain("logreg", bundle["clf"]["logreg"], X, bundle["features"])
    p = bundle["clf"]["logreg"].predict_proba(X)[0, 1]
    assert 1 / (1 + np.exp(-(sum(e["all"].values()) + e["intercept"]))) == pytest.approx(p, rel=1e-6)
    ex = explain.explain("xgboost", bundle["clf"]["xgboost"], X, bundle["features"])
    px = bundle["clf"]["xgboost"].predict_proba(X)[0, 1]
    assert 1 / (1 + np.exp(-(sum(ex["all"].values()) + ex["base_value"]))) == pytest.approx(px, rel=1e-4)


def test_weekly_records_pregame_metadata(week_run):
    t = week_run["table"]
    for c in ("predicted_at", "model_version", "predicted_before_kickoff", "feature_cutoff_utc", "home_projected_qb", "home_qb_uncertain"):
        assert c in t.columns
    assert (pd.to_datetime(t["feature_cutoff_utc"]) == pd.to_datetime(t["kickoff_utc"])).all()
    assert "freshness" in json.loads(week_run["paths"]["snapshot_json"].read_text())


def test_snapshots_are_immutable(week_run):
    p = week_run["paths"]["snapshot_json"]
    payload = json.loads(p.read_text())
    assert payload["meta"]["immutable"] and len(payload["predictions"]) == 4
    with pytest.raises(FileExistsError):
        open(p, "x").close()
    second = weekly.run_weekly(season=2026, week=6, refresh=False, progress=lambda m: None)
    assert second["paths"]["snapshot_json"] != p and p.exists()
    assert json.loads(p.read_text()) == payload


def test_excel_workbooks(week_run, trained):
    wb = load_workbook(week_run["paths"]["excel"])
    assert wb.sheetnames == ["Weekly Dashboard", "Game Predictions", "Six-Model Voting", "Model Probabilities", "Model Debate",
                             "Predicted Margins", "Injury & QB Monitor", "Upset Signals", "Experimental Adjusted",
                             "Sportsbook Comparison", "Model Information"]
    gp = wb["Game Predictions"]
    assert gp.freeze_panes == "B2" and gp.auto_filter.ref
    assert gp.max_row == 5
    from lab.reports import training_workbook
    tw = load_workbook(training_workbook())
    assert tw.sheetnames == ["Dashboard", "Model Leaderboard", "Historical Predictions", "Season Results", "Calibration", "Model Votes",
                             "Feature Importance", "Margin Model Results", "Sportsbook Comparison", "Data Coverage", "Methodology"]
    assert tw["Historical Predictions"].max_row > 100


@pytest.mark.parametrize("page", ["Home", "Data Manager", "Train Models", "Model Leaderboard", "Model Debate",
                                  "Weekly Predictions", "Upset Discovery", "Upset Explorer", "Research Validation",
                                  "Research Runs", "Reports"])
def test_streamlit_pages_start(page, week_run):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app" / "streamlit_app.py"), default_timeout=120)
    at.run()
    from app.navigation import SECTIONS
    section = next(section for section, pages in SECTIONS.items() if page in pages)
    at.sidebar.selectbox[0].select(section).run()
    at.sidebar.radio[0].set_value(page).run()
    assert at.sidebar.radio[0].value == page
    assert not at.exception, [e.value for e in at.exception]
    assert not at.error, [e.value for e in at.error]
    assert at.title[0].value
