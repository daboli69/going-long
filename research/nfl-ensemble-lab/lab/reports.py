"""Excel (openpyxl), CSV and JSON exports. Every number comes from a saved
training run or prediction snapshot."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import ensemble as ens
from . import models
from .config import REPORTS_DIR
from .data_pipeline import load_status, load_validation
from .features import feature_catalog
from .training import load_run

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F3A5F")
PROB_HINTS = ("p_", "prob", "accuracy", "observed", "mean_pred", "rate", "agreement", "share", "weight")


def _is_prob(col: str) -> bool:
    c = col.lower()
    return any(c.startswith(h) or h in c for h in PROB_HINTS) and "margin" not in c and "log_loss" not in c


def _clean(v):
    if isinstance(v, (list, tuple, dict)):
        return json.dumps(v, default=str)
    if isinstance(v, pd.Timestamp):
        return v.tz_convert(None).to_pydatetime() if v.tzinfo else v.to_pydatetime()
    if isinstance(v, (np.floating, float)) and np.isnan(v):
        return None
    if isinstance(v, np.generic):
        return v.item()
    return v


def write_sheet(wb: Workbook, title: str, df: pd.DataFrame, note: str | None = None, color_cols=()):
    ws = wb.create_sheet(title[:31])
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="555555")
        start = 3
    if df is None or df.empty:
        ws.cell(row=start, column=1, value="No data available for this sheet (nothing was fabricated).")
        return ws
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start, column=j, value=str(col))
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start + 1):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=_clean(v))
    last = start + len(df)
    for j, col in enumerate(df.columns, 1):
        letter = get_column_letter(j)
        width = max(len(str(col)), *(len(str(_clean(v))) for v in df[col].head(200))) if len(df) else len(str(col))
        ws.column_dimensions[letter].width = min(max(width + 2, 9), 48)
        if pd.api.types.is_float_dtype(df[col]):
            fmt = "0.0%" if _is_prob(str(col)) and df[col].dropna().between(0, 1).all() else "0.000"
            for r in range(start + 1, last + 1):
                ws.cell(row=r, column=j).number_format = fmt
        if col in color_cols:
            ws.conditional_formatting.add(f"{letter}{start + 1}:{letter}{last}", ColorScaleRule(
                start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=0.5,
                mid_color="FFFFFF", end_type="num", end_value=1, end_color="63BE7B"))
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(df.columns))}{last}"
    return ws


def _kv_sheet(wb, title, heading, items: list[tuple[str, object]]):
    ws = wb.create_sheet(title[:31])
    ws.cell(row=1, column=1, value=heading).font = TITLE_FONT
    for i, (k, v) in enumerate(items, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=_clean(v) if not isinstance(v, str) else v)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 110
    return ws


def _ci(v):
    if isinstance(v, str):
        try:
            v = json.loads(v.replace("(", "[").replace(")", "]"))
        except Exception:
            return v
    if isinstance(v, (list, tuple)) and len(v) == 2:
        return f"{v[0]:.3f} to {v[1]:.3f}"
    return v


METHODOLOGY = [
    ("Purpose", "Research lab only. Not a betting recommendation system; not connected to GOING production."),
    ("Data", "nflverse/nfldata games.csv (schedule, results, starting QBs, closing lines) + nflverse-data play-by-play (EPA, success)."),
    ("Leakage control", "Each feature uses only games whose kickoff is strictly earlier (as-of join). Ratings use games before the week's first kickoff."),
    ("Split", "Tune on 2016-2022 (folds 2020/2021/2022); select calibration + champion on 2023-2024; 2025 untouched test; 2026 prospective."),
    ("Walk-forward", "Each season s is predicted by models fit on all completed seasons < s. Development-season rows 2020-2022 share data with tuning, so they are optimistic."),
    ("Calibration", "Per model: none / Platt / isotonic chosen on validation; each season's calibrator is fit only on earlier out-of-sample seasons."),
    ("Ensembles", "Majority vote (3-3 broken by equal average), equal average, validation-skill-weighted average, out-of-fold logistic stacking."),
    ("Correlated votes", "Logistic Regression and the Single-Layer NN are the same model family. Six votes are not six independent opinions; 6-0 is not 100%."),
    ("Market lines", "nfldata closing consensus lines carry no timestamps. They are a benchmark only, never a feature."),
    ("Margins", "Point-margin regressors estimate expected home margin, not a final score."),
    ("Ties", "Tied games are excluded from winner metrics and included in margin metrics."),
    ("Profitability", "Winner accuracy and log loss say nothing about betting profit. No return backtest was run."),
]


def training_workbook(version: str | None = None) -> Path:
    run = load_run(version)
    if not run:
        raise FileNotFoundError("No training run found.")
    s, lb, oof = run["summary"], run["leaderboard"].copy(), run["oof"]
    for c in ("accuracy_ci", "brier_ci", "log_loss_ci"):
        if c in lb:
            lb[c] = lb[c].map(_ci)
    wb = Workbook()
    wb.remove(wb.active)
    test_ch = lb[(lb["split"] == "test") & (lb["method"] == s["champion"])]
    _kv_sheet(wb, "Dashboard", "NFL Ensemble Lab: Training Results", [
        ("Model version", s["version"]), ("Created (UTC)", s["created_at"]), ("Code commit", s.get("git_sha")),
        ("Completed games used", s["n_games"]), ("Seasons", f"{s['seasons'][0]}-{s['seasons'][1]}"),
        ("Features", s["n_features"]), ("Research champion (chosen on validation)", ens.MODEL_NAMES.get(s["champion"], ens.ENSEMBLE_NAMES.get(s["champion"], s["champion"]))),
        ("Champion 2025 test accuracy", float(test_ch["accuracy"].iloc[0]) if len(test_ch) else None),
        ("Champion 2025 test log loss", float(test_ch["log_loss"].iloc[0]) if len(test_ch) else None),
        ("Calibration chosen", json.dumps(s["calibration"])), ("Ensemble weights", json.dumps({k: round(v, 3) for k, v in s["weights"].items()})),
        ("Training minutes", s["minutes"]), ("Warning", "Historical accuracy is not evidence of betting profitability."),
    ])
    write_sheet(wb, "Model Leaderboard", lb, "Out-of-sample metrics by split. Champion was selected on validation before test was scored.", ["accuracy"])
    keep = ["game_id", "season", "week", "split", "home_team", "away_team", "home_win", "margin"] + \
           [c for c in oof.columns if c.startswith("p_")] + ["home_votes", "away_votes", "vote_tie", "margin_ensemble", "spread_line"]
    hp = oof[[c for c in keep if c in oof.columns]].copy()
    write_sheet(wb, "Historical Predictions", hp, "Walk-forward out-of-sample probabilities (home win). p_market = closing no-vig benchmark.",
                [c for c in hp.columns if c.startswith("p_")])
    write_sheet(wb, "Season Results", run["seasons"], None, ["accuracy"])
    write_sheet(wb, "Calibration", run["calibration"], "Predicted vs observed home-win rate by probability bucket.")
    votes = oof.groupby(["split", "home_votes", "away_votes"]).agg(games=("game_id", "size"), home_win_rate=("home_win", "mean")).reset_index()
    write_sheet(wb, "Model Votes", votes, "How often each vote split occurred and the observed home-win rate. Vote counts are not probabilities.")
    imp = run["importance"].sort_values(["model", "importance"], ascending=[True, False]).groupby("model").head(15) if len(run["importance"]) else run["importance"]
    write_sheet(wb, "Feature Importance", imp, "GLOBAL importance (permutation, out-of-sample season). Game-specific explanations are in weekly reports.")
    write_sheet(wb, "Margin Model Results", run["margins"].assign(mae_ci=run["margins"].get("mae_ci", pd.Series(dtype=object)).map(_ci)) if len(run["margins"]) else run["margins"])
    comp = lb[lb["method"].isin([s["champion"], "market", "base_pointdiff", "base_home_prior"])]
    write_sheet(wb, "Sportsbook Comparison", comp, "Closing no-vig lines are untimestamped consensus closes; benchmark only.")
    cov = pd.DataFrame(load_validation().get("coverage", []))
    write_sheet(wb, "Data Coverage", cov)
    _kv_sheet(wb, "Methodology", "Methodology and limitations", METHODOLOGY)
    path = REPORTS_DIR / "NFL_Ensemble_Training_Results.xlsx"
    wb.save(path)
    lb.to_csv(REPORTS_DIR / "training_leaderboard.csv", index=False)
    (REPORTS_DIR / "training_summary.json").write_text(json.dumps(s, indent=2, default=str), encoding="utf-8")
    return path


def weekly_workbook(season: int, week: int, table: pd.DataFrame, debates: list[dict], bundle: dict, run: dict,
                    monitor: list[dict] | None = None, freshness: dict | None = None, disc: dict | None = None) -> dict:
    keys = bundle["model_keys"]
    t = table.copy()
    wb = Workbook()
    wb.remove(wb.active)
    _kv_sheet(wb, "Weekly Dashboard", f"NFL Weekly Predictions: {season} Week {week}", [
        ("Generated (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")), ("Model version", bundle["version"]),
        ("Models trained through", bundle["trained_through"]), ("Games", len(t)),
        ("Research champion", bundle["champion"]), ("Unanimous votes", int(t["unanimous"].sum())),
        ("3-3 ties", int(t["vote_tie"].sum())), ("Vote vs ensemble disagreements", int(t["vote_vs_ensemble_disagree"].sum())),
        ("Games with data warnings", int((t["data_warnings"] != "").sum())),
        ("All predictions made before kickoff", bool(t["predicted_before_kickoff"].all()) if "predicted_before_kickoff" in t else "unknown"),
        ("Data refreshed (UTC)", (freshness or {}).get("data_refreshed_at")),
        ("Official injury file downloaded (UTC)", (freshness or {}).get("live_injury_file_downloaded_at")),
        ("ESPN injury feed retrieved (UTC)", (freshness or {}).get("espn_retrieved_at")),
        ("Reminder", "Probabilities are research estimates. A 6-0 vote is not a 100% forecast."),
    ])
    gp = pd.DataFrame({
        "game": t["away_team"] + " @ " + t["home_team"], "kickoff_utc": t["kickoff_utc"].astype(str),
        "predicted_before_kickoff": t.get("predicted_before_kickoff", pd.Series(dtype=bool)),
        "away_projected_qb": t.get("away_projected_qb", pd.Series(dtype=object)),
        "home_projected_qb": t.get("home_projected_qb", pd.Series(dtype=object)),
        "qb_uncertain": t.get("home_qb_uncertain", pd.Series(False, index=t.index)) | t.get("away_qb_uncertain", pd.Series(False, index=t.index)),
        "final_injury_reports": t.get("home_injury_report_final", pd.Series(False, index=t.index)) & t.get("away_injury_report_final", pd.Series(False, index=t.index)),
        **{ens.MODEL_NAMES[k]: np.where(t[f"p_{k}"] > 0.5, t["home_team"], t["away_team"]) for k in keys},
        "vote": t["vote_label"], "majority_winner": t["majority_winner"], "ensemble_winner": t["ensemble_winner"],
        "p_ensemble_home": t["p_champion"], "predicted_home_margin": t["margin_ensemble"].round(1),
        "model_disagreement_range": t["prob_spread"], "upset_research_signal": "Phase 2 not built: no signals evaluated",
        "data_warnings": t["data_warnings"],
    })
    write_sheet(wb, "Game Predictions", gp, None, ["p_ensemble_home"])
    write_sheet(wb, "Six-Model Voting", t[["game_id", "home_team", "away_team", "home_votes", "away_votes", "vote_label", "vote_tie", "unanimous", "majority_winner", "ensemble_winner", "vote_vs_ensemble_disagree"]])
    pcols = [f"p_{k}" for k in keys] + ["p_equal_average", "p_weighted_average", "p_stacking", "p_champion"]
    write_sheet(wb, "Model Probabilities", t[["game_id", "home_team", "away_team"] + pcols], "Calibrated home-win probabilities.", pcols)
    deb = []
    for d in debates:
        for r in d["round1"]:
            e = d["round2"][r["model"]]
            deb.append({"game": d["matchup"], "model": r["name"], "pick": r["winner"], "p_home": r["p_home"], "method": e["method"], "unit": e["unit"],
                        "top_reasons": "; ".join(f"{x['feature']}={x['value']:.3g} ({x['contribution']:+.3f} -> {x['favors']})" for x in e["top"][:5]),
                        "dissent": r["name"] in d["round3"]["dissenters"]})
    write_sheet(wb, "Model Debate", pd.DataFrame(deb), "Local (this-game) attributions computed from each fitted model.")
    mcols = [f"margin_{k}" for k in keys] + ["margin_ensemble"]
    pm = t[["game_id", "home_team", "away_team"] + mcols + ["spread_line"]].rename(columns={f"margin_{k}": models.MARGIN_NAMES[k] for k in keys})
    write_sheet(wb, "Predicted Margins", pm, "Expected home margin (points). spread_line = current line at data refresh, home perspective.")
    mon = pd.DataFrame([{**{k: v for k, v in m.items() if k not in ("key_listed_players", "espn_out_or_ir", "qb_flags")},
                         "qb_flags": "; ".join(m["qb_flags"]),
                         "key_listed_players": "; ".join(f"{x['player']} ({x['pos']}, {x['status']}, snaps {x['recent_snap_share']:.0%})" for x in m["key_listed_players"]),
                         "espn_out_or_ir": "; ".join(f"{x['player']} ({x['position']}, {x['status']})" for x in m["espn_out_or_ir"])}
                        for m in (monitor or [])])
    write_sheet(wb, "Injury & QB Monitor", mon, "LIVE MONITORING ONLY. ESPN and depth-chart data are not model inputs. "
                "Model injury features use only the official report available before kickoff.")
    disc = disc or {}
    sig_rows = []
    for r, sigs in zip(disc.get("rows", []), disc.get("signals", [])):
        for x in sigs:
            sig_rows.append({"game": f"{r['away_team']} @ {r['home_team']}", "underdog": r["underdog"], "pattern": x["pattern"],
                             "status": x["status"], "logodds_contribution": x.get("logodds_contribution"),
                             "missing_features": ", ".join(x.get("missing", []))})
    if disc.get("rows") and not sig_rows:
        sig_rows = [{"game": "all games", "pattern": "none", "status": "No validated discovery applies to any game this week "
                     "(the current experiment validated none)."}]
    write_sheet(wb, "Upset Signals", pd.DataFrame(sig_rows), "Patterns from the Upset Discovery Engine (experiment "
                f"{disc.get('experiment')}). Only VALIDATED patterns change probabilities; exploratory matches are context only.")
    ea = pd.DataFrame(disc.get("rows", []))
    if len(ea):
        ea = ea[["home_team", "away_team", "underdog", "p_home_original", "p_home_adj_patterns", "adj_patterns", "p_home_adj_gbm",
                 "adj_gbm", "n_validated_signals", "missing_gbm_inputs", "patterns_verdict", "gbm_verdict", "primary_forecast"]]
    write_sheet(wb, "Experimental Adjusted", ea, "EXPERIMENTAL. The original ensemble probability is the primary forecast unless "
                "a verdict reads 'demonstrated out-of-sample improvement'.", ["p_home_original", "p_home_adj_patterns", "p_home_adj_gbm"])
    sb = t[["game_id", "home_team", "away_team", "p_champion", "market_p_home_novig", "home_moneyline", "away_moneyline", "spread_line", "margin_ensemble"]].copy()
    sb["prob_minus_market"] = sb["p_champion"] - sb["market_p_home_novig"]
    write_sheet(wb, "Sportsbook Comparison", sb, "Lines are the nfldata line at refresh time (not timestamped by the source). No bet recommendations.")
    lb = run.get("leaderboard", pd.DataFrame())
    info = lb[lb["split"].isin(["validation", "test"])][["split", "name", "n", "accuracy", "brier", "log_loss", "ece"]] if len(lb) else lb
    write_sheet(wb, "Model Information", info, "Historical out-of-sample performance of each model and ensemble.")
    stem = f"NFL_Weekly_Predictions_{season}_{week:02d}"
    path = REPORTS_DIR / f"{stem}.xlsx"
    wb.save(path)
    csv = REPORTS_DIR / f"{stem}.csv"
    t.to_csv(csv, index=False)
    js = REPORTS_DIR / f"{stem}.json"
    js.write_text(json.dumps({"predictions": json.loads(t.to_json(orient="records", date_format="iso")), "debates": debates}, indent=1, default=str), encoding="utf-8")
    return {"excel": path, "csv": csv, "json": js}


def feature_catalog_csv(fs_columns) -> Path:
    p = REPORTS_DIR / "feature_catalog.csv"
    feature_catalog(fs_columns).to_csv(p, index=False)
    return p


DISCOVERY_METHODOLOGY = [
    ("What this is", "An automated search for situations where underdogs win more or less often than a baseline probability "
                     "says. It is research, not betting advice."),
    ("Baselines", "ensemble = the Phase 1 champion's out-of-sample probability; market = closing no-vig line (benchmark only: "
                  "no timestamps, not executable, so no profit claim is possible)."),
    ("Why one upset proves nothing", "A 25% underdog should win about 1 in 4 times. Patterns are judged on many games using "
                                     "likelihood statistics (observed wins minus expected wins), not on memorable results."),
    ("Periods", "Discovery seasons search for patterns; separate later validation seasons re-test the top 30 "
                "(Holm-corrected); a still-later evaluation season tests any probability adjustment. 2025 is diagnostic only; "
                "2026 counts only from predictions saved before kickoff."),
    ("Methods", "Threshold subgroups (single and paired), miscalibration buckets, RuleFit-style tree rules, symbolic "
                "expressions, gradient-boosted residual model with SHAP interactions, k-means regimes and anomaly detection."),
    ("Nominal vs valid significance", "Discovery-stage p/q values are nominal because the search adapts to the same games. "
                                      "Only the validation test counts."),
    ("Missing data", "Injury features are missing for 2025 and early 2026; a pattern needing them is reported as not "
                     "evaluable, never as false."),
    ("Adjustments", "Only validated patterns can adjust probabilities, through a shrunken (ridge) correction. The "
                    "gradient-boosted adjuster is scaled by a factor chosen on validation, so 0 means no adjustment. The original "
                    "ensemble stays primary unless an adjustment beats it on both evaluation seasons with a confidence interval "
                    "below zero."),
    ("Synthetic tests", "Planted-signal and null-data tests check the software only; they are not NFL findings."),
]


def discovery_workbook(exp: str | None = None, weekly_disc: dict | None = None) -> Path:
    from . import discovery as D
    e = D.load_experiment(exp)
    if not e:
        raise FileNotFoundError("No discovery experiment yet. Run one from Research Runs.")
    s, d = e["summary"], e["dir"]
    wb = Workbook()
    wb.remove(wb.active)
    items = [("Experiment", s["experiment"]), ("Created (UTC)", s["created_at"]), ("Runtime (minutes)", s.get("minutes"))]
    for bl, info in s["baselines"].items():
        for o, v in info["origins"].items():
            ev = v["evals"].get("evaluation", {})
            items.append((f"{bl} / {o}",
                          f"tested {v['n_tested']} candidates over {v['n_features']} features; {v['n_carried']} carried to "
                          f"validation; {v['n_validated']} validated. Evaluation seasons {ev.get('seasons')}: baseline log loss "
                          f"{ev.get('baseline', {}).get('log_loss', float('nan')):.4f}, patterns delta "
                          f"{ev.get('validated_patterns', {}).get('delta_log_loss', float('nan')):+.4f}, GBM delta "
                          f"{ev.get('gbm_residual', {}).get('delta_log_loss', float('nan')):+.4f}"))
        items.append((f"{bl} verdict", json.dumps(info["verdict"])))
        items.append((f"{bl} inputs", f"model {info['model_version']} (champion {info['champion']}), {info['rows']} games, "
                                      f"input hash {info['input_hash']}"))
    items.append(("Bottom line", "A pattern is only trusted if it survives independent validation. Zero validated patterns "
                                 "is a valid, honest result. Nothing here is a betting edge."))
    _kv_sheet(wb, "Plain-English Summary", "NFL Upset Discovery: what was found", items)
    cats = [c.assign(baseline=k.split("_", 1)[0], origin=k.split("_", 1)[1]) for k, c in e["catalogs"].items()]
    cat = pd.concat(cats, ignore_index=True) if cats else pd.DataFrame()
    lead_cols = ["baseline", "origin", "status", "method", "description", "disc_n", "disc_underdog_win_rate",
                 "disc_baseline_underdog_prob", "val_n", "val_underdog_win_rate", "val_baseline_underdog_prob",
                 "val_logodds_effect", "val_effect_ci_low", "val_effect_ci_high", "val_effect_shrunk", "val_p_holm",
                 "val_season_sign_share", "season_effects", "rejection_reasons"]
    carried = cat[cat["carried_to_validation"].astype(bool)] if len(cat) else cat
    write_sheet(wb, "Discovery Leaderboard", carried[[c for c in lead_cols if c in carried]] if len(carried) else carried,
                "Ranked by validated, shrunk out-of-sample effect. rejected = failed independent validation.")
    rule_cols = ["baseline", "origin", "status", "method", "description", "expression", "features", "provenance",
                 "disc_n", "disc_z", "disc_p_nominal", "disc_q_bh_nominal"]
    write_sheet(wb, "Discovered Rules", cat[[c for c in rule_cols if c in cat]] if len(cat) else cat,
                "Every candidate carried to validation plus nominally notable exploratory ones (discovery q < 0.5). "
                "Discovery p/q values are nominal.")
    write_sheet(wb, "Feature Interactions",
                pd.concat([v.assign(source=k) for k, v in e["interactions"].items()], ignore_index=True)
                if e["interactions"] else pd.DataFrame(),
                "Strongest pairwise interactions in the discovery-period residual model (SHAP).")
    m = [pd.read_csv(f).assign(source=f.stem[8:]) for f in d.glob("matches_*.csv") if f.stat().st_size > 5]
    write_sheet(wb, "Historical Matches", pd.concat(m, ignore_index=True) if m else pd.DataFrame(),
                "Games that matched validated patterns and the top carried patterns, with outcomes.")
    cal = [pd.read_csv(f).assign(baseline=f.stem[12:]) for f in d.glob("calibration_*.csv")]
    write_sheet(wb, "Expected vs Observed", pd.concat(cal, ignore_index=True) if cal else pd.DataFrame(),
                "Underdog win rate expected by the baseline vs observed, by probability bucket and season.")
    ev_rows = []
    for bl, info in s["baselines"].items():
        for o, v in info["origins"].items():
            for period, ev in v["evals"].items():
                for layer in ("baseline", "validated_patterns", "gbm_residual"):
                    x = ev.get(layer, {})
                    ev_rows.append({"baseline": bl, "origin": o, "period": period, "seasons": str(ev.get("seasons")),
                                    "layer": layer, "n": x.get("n"), "log_loss": x.get("log_loss"), "brier": x.get("brier"),
                                    "ece": x.get("ece"), "delta_log_loss_vs_baseline": x.get("delta_log_loss"),
                                    "delta_ci": str(x.get("delta_ci", "")), "mean_abs_adjustment": x.get("mean_abs_adjustment")})
    write_sheet(wb, "Validation Results", pd.DataFrame(ev_rows), "Negative delta = adjusted forecast beat the baseline. "
                "Evaluation seasons were not used for discovery or tuning.")
    rej = carried[carried["status"] == "rejected"] if len(carried) else carried
    rej_cols = ["baseline", "origin", "method", "description", "disc_n", "disc_z", "val_n", "val_logodds_effect",
                "val_p_holm", "rejection_reasons"]
    write_sheet(wb, "Rejected Hypotheses", rej[[c for c in rej_cols if c in rej]] if len(rej) else rej)
    write_sheet(wb, "Upcoming Upset Candidates", pd.DataFrame((weekly_disc or {}).get("rows", [])),
                "Upcoming games with original and experimental probabilities. No game is labelled an upset pick.")
    _kv_sheet(wb, "Methodology", "How the discovery engine works, and its limits", DISCOVERY_METHODOLOGY)
    path = REPORTS_DIR / "NFL_Upset_Discovery.xlsx"
    wb.save(path)
    if len(cat):
        cat.to_csv(REPORTS_DIR / "upset_discovery_catalog.csv", index=False)
    (REPORTS_DIR / "upset_discovery_summary.json").write_text(json.dumps(s, indent=2, default=str), encoding="utf-8")
    return path
